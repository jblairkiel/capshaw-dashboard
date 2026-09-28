"""A sample order of service as a Word document, for This Sunday to show.
The people in it come from the run's made-up directory.

    python3 setup/order_of_service.py <work dir>   ->  <work dir>/assets/order-of-service.docx
"""
import json, os, sqlite3, sys
from datetime import date, timedelta
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

work = sys.argv[1]
facts = json.load(open(os.path.join(work, 'facts.json')))
db = sqlite3.connect(os.path.join(work, 'data', 'bible_questions.db'))
names = [n for (n,) in db.execute("SELECT name FROM directory WHERE gender = 'male' ORDER BY id")]
names = [n for n in names if n != facts['me']['name']] or names
who = lambda i: names[i % len(names)]

sunday = date.today() - timedelta(days=(date.today().weekday() + 1) % 7)
doc = Document()
doc.add_heading('Order of Worship', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(f'Sunday Morning · {sunday:%B} {sunday.day}, {sunday.year}').alignment = WD_ALIGN_PARAGRAPH.CENTER

rows = [
    ('Welcome & Announcements', who(0)),
    ('Song #916 — Night With Ebon Pinion', f'{who(1)} leading'),
    ('Opening Prayer', facts['me']['name']),
    ('Song #183 — Hallelujah Praise Jehovah', ''),
    ('Scripture Reading — Galatians 6:7-10', who(2)),
    ('Song #774 — Sing to Me of Heaven', ''),
    ("The Lord's Supper", who(3)),
    ('Sermon — "Faith That Works"', who(4)),
    ('Song of Invitation #726 — Just As I Am', ''),
    ('Closing Prayer', who(5)),
]
table = doc.add_table(rows=0, cols=2)
table.style = 'Light List Accent 1'
for part, person in rows:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = part, person
doc.add_paragraph()
doc.add_paragraph('Fellowship meal after evening services — everyone welcome.')

os.makedirs(os.path.join(work, 'assets'), exist_ok=True)
doc.save(os.path.join(work, 'assets', 'order-of-service.docx'))
print('  order of service written')
